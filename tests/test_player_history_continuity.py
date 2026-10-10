"""Player-history roster-lineage continuity guard.

A player's rendered history (the hover-comment chain on player_all_time / picks)
must never skip a step: a player can only be dropped or traded away by a team
that holds them, and can only be added off waivers/free agency while no team
holds them. A break means the player teleported on/off a roster — the class of
bug fixed by the history-ordering + missing-departure-synthesis passes in
lotg.py.

This test reuses the same reconstruction as scripts/audit_player_history.py and
asserts the freshly built workbook has no continuity breaks. It is skipped when
no workbook is present (e.g. a CSV-only check), so it never fails a partial run.

The workbook is a build output and is no longer committed (see .gitignore), so
a bare checkout skips this. That costs nothing in CI, which is where the guard
has to hold: build.yml builds before it runs pytest, with LOTG_EXPORTS=exports,
so the workbook is on disk and this asserts against it exactly as before.
"""
from __future__ import annotations

import os
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]


def _xlsx_path() -> Path:
    return Path(os.environ.get("LOTG_EXPORTS", REPO / "exports")) / "LOTG_Stats.xlsx"


def _is_fixera_build() -> bool:
    # The roster-lineage reconciliation logs this marker on every fix-era build.
    # Its absence means the workbook predates the fix; skip rather than assert
    # against pre-fix histories. CI builds fresh before pytest, so the marker is
    # present and the test runs.
    log = Path(os.environ.get("LOTG_EXPORTS", REPO / "exports")) / "raw" / "build_debug.log"
    try:
        return "orphaned roster lineage" in log.read_text(errors="ignore")
    except Exception:
        return False


@pytest.mark.skipif(not _xlsx_path().exists() or not _is_fixera_build(),
                    reason="no fix-era workbook to audit")
def test_no_player_history_continuity_breaks():
    pytest.importorskip("openpyxl")
    import sys

    sys.path.insert(0, str(REPO / "scripts"))
    import audit_player_history as aud

    comments = aud.load_history_comments(_xlsx_path())
    players = {}
    for k, v in comments.items():
        sheet, name = k.split(":", 1)
        if sheet == "player_all_time":
            players[name] = v
    for k, v in comments.items():
        sheet, name = k.split(":", 1)
        if sheet == "picks":
            players.setdefault(name, v)

    breaks = []
    for name, txt in players.items():
        breaks.extend(b for b in aud.audit_text(name, txt) if b[2] != "UNPARSED")

    detail = "\n".join(f"  {b[0]} {b[1]} {b[2]}: {b[4]}" for b in sorted(breaks)[:60])
    assert not breaks, f"{len(breaks)} player-history continuity break(s):\n{detail}"


def test_the_xml_comment_reader_matches_openpyxl(tmp_path=None):
    """`load_history_comments` reads the comments straight from the xlsx XML
    (openpyxl.load_workbook took ~2.5 of CI's test minutes). It must answer
    exactly as the openpyxl reader it replaced: str(value) keys (numbers cast
    as openpyxl casts them), the first row kept, column A only, and rich-text
    comments (the build bolds verbs into runs) joined to plain text."""
    openpyxl = pytest.importorskip("openpyxl")
    import sys
    import tempfile
    from openpyxl.comments import Comment

    sys.path.insert(0, str(REPO / "scripts"))
    sys.path.insert(0, str(REPO / "src"))
    import audit_player_history as aud
    from lotg import _bold_comment_verbs

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pa = wb.create_sheet("player_all_time")
    pa.append(["Player", "Points"])
    rows = [("Ja'Marr Chase", "2021-04-29: drafted by plehv79 (1.01)\n2023-01-02: traded to AceMatthew (x)"),
            ("Tyler Conklin", "2022-09-01: added by LWebs53 (free agent)\n2022-10-01: dropped by LWebs53"),
            ("Ja'Marr Chase", "a second row with the same name — the first one wins"),
            (None, "a comment on an empty name cell")]
    for i, (name, text) in enumerate(rows, start=2):
        pa.append([name, 1.5])
        pa.cell(row=i, column=1).comment = Comment(text, "LOTG")
    pa.cell(row=2, column=2).comment = Comment("column B — never read", "LOTG")
    rk = wb.create_sheet("rookie_picks")
    rk.append(["Year", "Number"])
    rk.append([2024, 3])
    rk.cell(row=2, column=1).comment = Comment("2024 rookie draft: plehv79 took Malik Nabers", "LOTG")
    wb.create_sheet("team_week").append(["Team"])
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "w.xlsx"
        wb.save(path)
        _bold_comment_verbs(path)              # rich-text runs, as the build writes them
        fast = aud.load_history_comments(path)
        slow = aud._load_history_comments_openpyxl(path)
    assert fast == slow, (fast, slow)
    assert fast["player_all_time:Ja'Marr Chase"].startswith("2021-04-29: drafted by")
    assert "rookie_picks:2024" in fast and "player_all_time:None" in fast
